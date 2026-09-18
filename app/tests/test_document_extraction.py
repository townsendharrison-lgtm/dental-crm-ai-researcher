import io
import json
import logging
from pathlib import Path
from unittest.mock import Mock

from docx import Document
import httpx
from openai import AsyncOpenAI
import pytest
from pydantic import ValidationError

from app.agents.document_extraction_agent import DocumentExtractionAgent
from app.clients.fetch_client import DocumentChunk, DocumentError, DocumentParser, OcrUnavailable, normalize_text
from app.clients.llm_client import LLMClient, ExtractionFailed, ModelRefused
from app.factor_taxonomy import FactorTaxonomy, TaxonomyNotConfigured, get_taxonomy
from app.schemas.school_profile import GroundingError, validate_result

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def taxonomy():
    # Test-only definitions. These do not establish the production business taxonomy.
    return FactorTaxonomy.model_validate({
        "version": "synthetic-test-only", "categories": ["Test numbers", "Test mission"],
        "factors": [
            {"key": "avg_gpa", "category": "Test numbers", "description": "Average overall GPA", "value_type": "number", "unit": "gpa"},
            {"key": "avg_science_gpa", "category": "Test numbers", "description": "Average science GPA", "value_type": "number", "unit": "gpa"},
            {"key": "avg_dat", "category": "Test numbers", "description": "Average DAT", "value_type": "number", "unit": "points"},
            {"key": "mission", "category": "Test mission", "description": "School mission", "value_type": "text", "unit": None},
        ],
    })


def fact(**changes):
    fields = {"factor_key": "avg_gpa", "value": 3.7, "unit": "gpa",
              "raw_text_snippet": "Average overall GPA 3.7", "page_number": 1, "confidence": 0.95}
    fields.update(changes)
    return fields


def response(facts, *, refusal=None, finish="stop", raw=None, usage=True):
    result = {"id": "fixture-completion", "object": "chat.completion", "created": 1, "model": "gpt-4o",
              "choices": [{"index": 0, "finish_reason": finish, "message": {
                  "role": "assistant", "content": raw if raw is not None else json.dumps({"facts": facts}),
                  "refusal": refusal,
              }}]}
    if usage:
        result["usage"] = {"prompt_tokens": 100, "completion_tokens": 20, "total_tokens": 120,
                           "prompt_tokens_details": {"cached_tokens": 40}}
    return result


def make_llm(settings, responses):
    calls = []
    iterator = iter(responses)
    async def handle(request):
        calls.append(json.loads(request.content))
        item = next(iterator)
        if isinstance(item, int):
            return httpx.Response(item, json={"error": {"message": "synthetic provider error", "type": "api_error"}})
        return httpx.Response(200, json=item)
    http_client = httpx.AsyncClient(transport=httpx.MockTransport(handle))
    sdk = AsyncOpenAI(api_key="synthetic-key", http_client=http_client, max_retries=0)
    return LLMClient(settings, client=sdk), calls


def test_pdf_page_and_table_text_are_preserved(settings):
    parser = DocumentParser(settings)
    chunks = parser.parse((FIXTURES / "class_profile.pdf").read_bytes(), "class_profile.pdf")
    assert len(chunks) == 1
    assert chunks[0].page_number == 1
    assert "Average overall GPA 3.7" in chunks[0].text
    assert "Average science GPA 3.6" in chunks[0].text
    assert not chunks[0].used_ocr


def test_docx_sections_and_tables_without_invented_pages(settings):
    doc = Document()
    doc.add_heading("Admissions", level=1)
    doc.add_paragraph("Average overall GPA 3.7")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Average science GPA", "3.6"
    doc.add_heading("Mission", level=1)
    doc.add_paragraph("Service to underserved communities")
    stream = io.BytesIO(); doc.save(stream)
    chunks = DocumentParser(settings).parse(stream.getvalue(), "test.docx")
    assert [c.section for c in chunks] == ["Admissions", "Mission"]
    assert all(c.page_number is None for c in chunks)
    assert "Average science GPA | 3.6" in chunks[0].text


def test_long_sections_are_not_truncated(settings):
    parser = DocumentParser(settings.model_copy(update={"document_chunk_chars": 500}))
    content = "\n".join(f"Line {i}: sample admissions text and evidence." for i in range(200))
    chunks = list(parser._split(content))
    assert len(chunks) > 10 and all(len(chunk) <= 500 for chunk in chunks)
    assert normalize_text(" ".join(chunks)) == normalize_text(content)


def test_scanned_pdf_uses_ocr_with_real_page_image(settings):
    ocr = Mock(return_value="Average overall GPA 3.7")
    parser = DocumentParser(settings, ocr=ocr)
    data = (FIXTURES / "scanned_profile.pdf").read_bytes()
    chunks = parser.parse(data, "scan.pdf")
    ocr.assert_called_once_with(data, 0)
    assert chunks[0].used_ocr and chunks[0].page_number == 1


def test_missing_tesseract_is_explicit_failure(settings):
    parser = DocumentParser(settings.model_copy(update={"tesseract_cmd": "definitely-missing-tesseract-fixture"}))
    with pytest.raises(OcrUnavailable):
        parser.parse((FIXTURES / "scanned_profile.pdf").read_bytes(), "scan.pdf")


@pytest.mark.parametrize("data,name", [(b"hello", "x.pdf"), (b"%PDF-fake", "x.exe"), (b"", "x.docx")])
def test_invalid_upload_is_rejected(settings, data, name):
    with pytest.raises(DocumentError):
        DocumentParser(settings).parse(data, name)


def test_size_limit_is_enforced(settings):
    parser = DocumentParser(settings.model_copy(update={"document_max_bytes": 20}))
    with pytest.raises(DocumentError):
        parser.parse((FIXTURES / "class_profile.pdf").read_bytes(), "profile.pdf")


@pytest.mark.parametrize("changes", [
    {"factor_key": "invented_factor"}, {"value": 3.8}, {"value": True},
    {"raw_text_snippet": "Fabricated evidence"}, {"page_number": 2},
    {"confidence": 1.1}, {"confidence": None}, {"unit": "percent"},
    {"value": None, "confidence": 0.9}, {"extra_field": "not allowed"},
])
def test_unsupported_or_malformed_facts_are_rejected(taxonomy, changes):
    chunk = DocumentChunk(0, "Average overall GPA 3.7", 1, None, False)
    with pytest.raises((GroundingError, ValidationError)):
        validate_result(json.dumps({"facts": [fact(**changes)]}), chunk, taxonomy)


async def test_realistic_pdf_extraction_and_missing_factor_omission(settings, taxonomy):
    llm, calls = make_llm(settings, [response([fact(), fact(factor_key="avg_science_gpa", value=3.6,
                                                         raw_text_snippet="Average science GPA 3.6")])])
    try:
        result = await DocumentExtractionAgent(DocumentParser(settings), llm, taxonomy).extract(
            (FIXTURES / "class_profile.pdf").read_bytes(), "class_profile.pdf")
    finally:
        await llm.close()
    assert result.complete
    facts = result.chunks[0].extraction.result.facts
    assert {f.factor_key: f.value for f in facts} == {"avg_gpa": 3.7, "avg_science_gpa": 3.6}
    assert "avg_dat" not in {f.factor_key for f in facts}
    assert calls[0]["model"] == "gpt-4o"
    schema = calls[0]["response_format"]["json_schema"]
    assert schema["strict"] is True
    assert schema["schema"]["$defs"]["ExtractedFact"]["properties"]["factor_key"]["enum"] == list(taxonomy.by_key)


async def test_mission_pdf_has_no_invented_numeric_facts(settings, taxonomy):
    quote = "Our mission is to improve oral health through service to underserved communities."
    llm, _ = make_llm(settings, [response([fact(factor_key="mission", value="service to underserved communities",
                                              unit=None, raw_text_snippet=quote)])])
    try:
        result = await DocumentExtractionAgent(DocumentParser(settings), llm, taxonomy).extract(
            (FIXTURES / "mission_values.pdf").read_bytes(), "mission.pdf")
    finally:
        await llm.close()
    assert result.complete
    assert [f.factor_key for f in result.chunks[0].extraction.result.facts] == ["mission"]


async def test_invalid_output_is_corrected_once(settings, taxonomy):
    llm, calls = make_llm(settings, [response([fact(value=99)]), response([fact()])])
    try:
        result = await llm.extract(DocumentChunk(0, "Average overall GPA 3.7", 1, None, False), taxonomy)
    finally:
        await llm.close()
    assert len(calls) == 2 and len(result.usage) == 2
    assert result.result.facts[0].value == 3.7


async def test_repeated_invalid_output_marks_chunk_failed(settings, taxonomy):
    llm, calls = make_llm(settings, [response([fact(value=99)]), response([fact(value=99)])])
    try:
        result = await DocumentExtractionAgent(DocumentParser(settings), llm, taxonomy).extract(
            (FIXTURES / "class_profile.pdf").read_bytes(), "class.pdf")
    finally:
        await llm.close()
    assert not result.complete and len(calls) == 2
    assert result.chunks[0].extraction is None
    assert result.chunks[0].error_type == "ExtractionFailed"


async def test_refusal_is_not_retried_or_invented(settings, taxonomy):
    llm, calls = make_llm(settings, [response([], refusal="fixture refusal")])
    try:
        with pytest.raises(ModelRefused):
            await llm.extract(DocumentChunk(0, "Average overall GPA 3.7", 1, None, False), taxonomy)
    finally:
        await llm.close()
    assert len(calls) == 1


async def test_transient_provider_error_uses_bounded_retry(settings, taxonomy):
    llm, calls = make_llm(settings, [429, response([])])
    try:
        result = await llm.extract(DocumentChunk(0, "No numeric information.", 1, None, False), taxonomy)
    finally:
        await llm.close()
    assert len(calls) == 2 and result.result.facts == []


async def test_usage_and_estimated_cost_are_logged_without_source_text(settings, taxonomy, caplog):
    llm, _ = make_llm(settings, [response([])])
    with caplog.at_level(logging.INFO, logger="school_ai.external"):
        try:
            await llm.extract(DocumentChunk(0, "PRIVATE_SOURCE_TEXT", 1, None, False), taxonomy)
        finally:
            await llm.close()
    record = json.loads(next(r.message for r in caplog.records if r.name == "school_ai.external"))
    assert record["tokens"] == 120 and record["cached_input_tokens"] == 40
    assert record["cost_usd"] == pytest.approx(0.0004)
    assert "PRIVATE_SOURCE_TEXT" not in caplog.text and "synthetic-key" not in caplog.text


def test_missing_production_taxonomy_fails_explicitly(monkeypatch):
    monkeypatch.setattr("app.factor_taxonomy.APPROVED_TAXONOMY", None)
    with pytest.raises(TaxonomyNotConfigured):
        get_taxonomy()


def test_numeric_evidence_allows_sentence_punctuation(taxonomy):
    chunk = DocumentChunk(0, "Average overall GPA 3.7.", 1, None, False)
    result = validate_result(json.dumps({"facts": [fact(raw_text_snippet=chunk.text)]}), chunk, taxonomy)
    assert result.facts[0].value == 3.7


async def test_truncated_response_is_not_accepted(settings, taxonomy):
    llm, calls = make_llm(settings, [response([fact()], finish="length"), response([fact()], finish="length")])
    try:
        with pytest.raises(ExtractionFailed):
            await llm.extract(DocumentChunk(0, "Average overall GPA 3.7", 1, None, False), taxonomy)
    finally:
        await llm.close()
    assert len(calls) == 2
